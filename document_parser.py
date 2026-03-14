"""Utilities for extracting rich content from PDF and PPTX documents."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from typing import Dict, List

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SlideRecord = Dict[str, object]


def _dedupe_images(images: List[bytes]) -> List[bytes]:
    unique: List[bytes] = []
    seen_hashes: set[str] = set()

    for blob in images:
        digest = hashlib.md5(blob).hexdigest()
        if digest not in seen_hashes:
            seen_hashes.add(digest)
            unique.append(blob)

    return unique


def extract_from_pdf(file_bytes: bytes) -> List[SlideRecord]:
    """Extract text, images, and page preview from a PDF file."""
    records: List[SlideRecord] = []

    with fitz.open(stream=file_bytes, filetype="pdf") as pdf_doc:
        for page_index, page in enumerate(pdf_doc):
            text = page.get_text("text").strip() or "[No readable text found on this page.]"

            # A lightweight page render used for preview and optional vision analysis.
            preview_pix = page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False)
            preview_bytes = preview_pix.tobytes("png")

            extracted_images: List[bytes] = []
            for image_info in page.get_images(full=True):
                xref = image_info[0]
                image_data = pdf_doc.extract_image(xref)
                image_bytes = image_data.get("image", b"")
                if image_bytes:
                    extracted_images.append(image_bytes)

            extracted_images = _dedupe_images(extracted_images)
            if not extracted_images:
                # Fallback: use page preview so visual layout can still be analyzed.
                extracted_images = [preview_bytes]

            records.append(
                {
                    "index": page_index,
                    "type": "pdf",
                    "text": text,
                    "preview_image": preview_bytes,
                    "images": extracted_images,
                }
            )

    return records


def extract_from_pptx(file_bytes: bytes) -> List[SlideRecord]:
    """Extract text and embedded images from a PPTX file."""
    records: List[SlideRecord] = []
    presentation = Presentation(BytesIO(file_bytes))

    for slide_index, slide in enumerate(presentation.slides):
        text_parts: List[str] = []
        images: List[bytes] = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                shape_text = shape.text.strip()
                if shape_text:
                    text_parts.append(shape_text)

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and hasattr(shape, "image"):
                image_blob = shape.image.blob
                if image_blob:
                    images.append(image_blob)

        images = _dedupe_images(images)
        combined_text = "\n".join(text_parts).strip() or "[No readable text found on this slide.]"
        preview_image = images[0] if images else None

        records.append(
            {
                "index": slide_index,
                "type": "pptx",
                "text": combined_text,
                "preview_image": preview_image,
                "images": images,
            }
        )

    return records


def parse_document(file_name: str, file_bytes: bytes) -> List[SlideRecord]:
    """Parse uploaded document and return one rich record per page/slide."""
    suffix = Path(file_name).suffix.lower()

    if suffix == ".pdf":
        return extract_from_pdf(file_bytes)

    if suffix == ".pptx":
        return extract_from_pptx(file_bytes)

    raise ValueError("Unsupported file format. Please upload a PDF or PPTX file.")
